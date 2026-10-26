"""Read-only source audit. Outputs are confined to research_audit/."""

from common import workspace_relative

from pathlib import Path
from bootstrap import ROOT
from common import safe_source, config
import os
from collections import Counter
import csv, json, re, zipfile
import openpyxl
import numpy as np
from PIL import Image, ExifTags

OUT = ROOT / "research_audit"


def write_csv(name, rows):
    if not rows:
        return
    fields = list(dict.fromkeys(k for row in rows for k in row))
    with (OUT / name).open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def source_inventory():
    """Inventory declared raw folders/workbooks, excluding execution artifacts."""
    files = {safe_source(path.relative_to(ROOT)) for path in ROOT.glob("*.xlsx")}
    folders = list(config().get("source_folders", []))
    if (ROOT / "about").is_dir():
        folders.append("about")
    skip = {"__pycache__", ".git"}
    visited = set()
    for folder in folders:
        source = safe_source(folder)
        if not source.is_dir():
            continue
        for directory, children, names in os.walk(source, followlinks=True):
            current = Path(directory)
            real = current.resolve()
            if real in visited:
                children[:] = []
                continue
            safe_source(current.relative_to(ROOT))
            visited.add(real)
            children[:] = [name for name in children if name not in skip]
            files.update(safe_source((current / name).relative_to(ROOT)) for name in names)
    return sorted(files)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    files = source_inventory()
    write_csv(
        "file_inventory.csv",
        [{"path": workspace_relative(p), "bytes": p.stat().st_size} for p in files],
    )
    tables = {}
    profiles = []
    diagnostics = []
    export = OUT / "tables"
    export.mkdir(exist_ok=True)
    for p in sorted(ROOT.glob("*.xlsx")):
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        for sheet in wb:
            rows = list(sheet.values)
            target = export / (p.stem + "__" + sheet.title + ".csv")
            with target.open("w", newline="", encoding="utf-8-sig") as f:
                csv.writer(f).writerows(rows)
            profiles.append(
                {
                    "file": p.name,
                    "sheet": sheet.title,
                    "rows_with_header": len(rows),
                    "columns": sheet.max_column,
                    "export": str(target.relative_to(OUT)),
                }
            )
            if sheet == wb.worksheets[0] and p.name not in [
                "Genotype_list.xlsx",
                "MASA_P_6.xlsx",
                "LAB_1.xlsx",
                "LAB_2.xlsx",
            ]:
                keys = [tuple(row[:4]) for row in rows[1:] if row[0] is not None]
                tables[p.name] = (rows, keys)
                diagnostics.append(
                    {
                        "file": p.name,
                        "records": len(keys),
                        "unique_keys": len(set(keys)),
                        "missing_numeric_cells": sum(
                            v is None for row in rows[1:] for v in row[4:]
                        ),
                        "excel_errors": sum(
                            isinstance(v, str)
                            and v.startswith(("#DIV/0!", "#VALUE!", "#REF!", "#N/A", "#NUM!"))
                            for row in rows[1:]
                            for v in row
                        ),
                    }
                )
        wb.close()
    write_csv("table_inventory.csv", profiles)
    write_csv("table_quality.csv", diagnostics)
    base = set(tables["OJIP_P6_computed_data.xlsx"][1])
    joins = [
        {
            "file": name,
            "matched_ojip_keys": len(base.intersection(keys)),
            "missing_vs_ojip": len(base - set(keys)),
            "extra_vs_ojip": len(set(keys) - base),
        }
        for name, (_, keys) in tables.items()
    ]
    write_csv("reference_key_alignment.csv", joins)
    genotype_rows = list(
        openpyxl.load_workbook(
            ROOT / "Genotype_list.xlsx", read_only=True, data_only=True
        ).active.values
    )
    genotype_map = {str(code): name for name, code in genotype_rows}
    imgs = []
    sessions = []
    session_paths = {
        p.parent.parent
        for p in files
        if p.suffix.lower() == ".jpg"
        and p.parent.name.lower() in ["rgb", "noir"]
        and p.parent.parent.name.startswith("capture_session")
    }
    for session in sorted(session_paths):
        match = re.fullmatch(r"capture_session_([CSP])(\d+)_(\d+)-(\d+)", session.name)
        key = None
        if match:
            t, m, g, rep = match.groups()
            key = (t, int(m), genotype_map["L" + g], int(rep))
        counts = Counter()
        good = Counter()
        indices = {"rgb": set(), "noir": set()}
        dates = set()
        exposures = set()
        isos = set()
        for camera in ["rgb", "noir"]:
            for p in sorted((session / camera).glob("*.jpg")):
                counts[camera] += 1
                error = ""
                width = height = None
                date = exposure = iso = None
                try:
                    with Image.open(p) as im:
                        width, height = im.size
                        ex = im.getexif().get_ifd(34665)
                        date = ex.get(36867)
                        exposure = float(ex[33434]) if 33434 in ex else None
                        iso = ex.get(34855)
                        im.verify()
                    good[camera] += 1
                except Exception as e:
                    error = type(e).__name__ + ": " + str(e)
                idx = int(p.stem.rsplit("_", 1)[1])
                indices[camera].add(idx)
                if date:
                    dates.add(date[:10])
                if exposure is not None:
                    exposures.add(exposure)
                if iso is not None:
                    isos.add(iso)
                imgs.append(
                    {
                        "path": workspace_relative(p),
                        "session": workspace_relative(session),
                        "camera": camera,
                        "view": idx,
                        "bytes": p.stat().st_size,
                        "width": width,
                        "height": height,
                        "datetime_original": date,
                        "exposure_s": exposure,
                        "iso": iso,
                        "container_check_error": error,
                    }
                )
        sessions.append(
            {
                "session": workspace_relative(session),
                "treatment": key[0] if key else "",
                "measurement": key[1] if key else "",
                "genotype": key[2] if key else "",
                "replicate": key[3] if key else "",
                "id_status": "filename_key_matches_reference"
                if key in base
                else "unresolved_or_pilot",
                "rgb_files": counts["rgb"],
                "noir_files": counts["noir"],
                "rgb_container_ok": good["rgb"],
                "noir_container_ok": good["noir"],
                "paired_view_indices": len(indices["rgb"] & indices["noir"]),
                "exif_dates": ";".join(sorted(dates)),
                "exposure_values": ";".join(map(str, sorted(exposures))),
                "iso_values": ";".join(map(str, sorted(isos))),
            }
        )
    write_csv("image_manifest.csv", imgs)
    write_csv("session_manifest.csv", sessions)
    rows, _ = tables["OJIP_P6_computed_data.xlsx"]
    headers = rows[0]
    fv = headers.index("Fv/Fm")
    stats = []
    for t in "CPS":
        for m in range(1, 6):
            values = np.array(
                [row[fv] for row in rows[1:] if row[0] == t and row[1] == m], dtype=float
            )
            stats.append(
                {
                    "treatment": t,
                    "measurement": m,
                    "n": len(values),
                    "Fv_Fm_mean": float(values.mean()),
                    "Fv_Fm_sd": float(values.std(ddof=1)),
                    "Fv_Fm_min": float(values.min()),
                    "Fv_Fm_max": float(values.max()),
                }
            )
    write_csv("ojip_descriptive_statistics.csv", stats)
    spectral_checks = []
    for prefix in ["UVIS", "NIR"]:
        ar, ak = tables[prefix + "_P6_absorbance.xlsx"]
        tr, tk = tables[prefix + "_P6_transmitance.xlsx"]
        a = {tuple(row[:4]): np.array(row[4:], float) for row in ar[1:]}
        tt = {tuple(row[:4]): np.array(row[4:], float) for row in tr[1:]}
        errors = []
        invalid = 0
        for key in a.keys() & tt.keys():
            valid = np.isfinite(tt[key]) & (tt[key] > 0) & np.isfinite(a[key])
            invalid += int((~valid).sum())
            errors.extend(np.abs(a[key][valid] + np.log10(tt[key][valid])).tolist())
        spectral_checks.append(
            {
                "instrument_label": prefix,
                "wavelength_min": ar[0][4],
                "wavelength_max": ar[0][-1],
                "channels": len(ar[0]) - 4,
                "A_vs_minus_log10_T_max_abs_error": max(errors),
                "invalid_pairs": invalid,
            }
        )
    write_csv("spectral_consistency.csv", spectral_checks)
    summary = {
        "source_files": len(files),
        "source_bytes": sum(p.stat().st_size for p in files),
        "jpg_count": len(imgs),
        "jpg_container_failures": sum(bool(x["container_check_error"]) for x in imgs),
        "sessions": len(sessions),
        "sessions_with_reference_key": sum(
            x["id_status"] == "filename_key_matches_reference" for x in sessions
        ),
        "reference_records": len(base),
        "reference_key_alignment": joins,
        "spectral_checks": spectral_checks,
    }
    (OUT / "audit_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
