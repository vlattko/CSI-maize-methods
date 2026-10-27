"""Rebuild keyed reference tables from unchanged acquisition workbooks."""

from __future__ import annotations

import csv
import hashlib
import json
import math
from pathlib import Path, PureWindowsPath
import re
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from PIL import Image

KEY_FIELDS = ("treatment", "measurement", "genotype", "replicate")
COMPUTED = (
    ("ojip", "OJIP_P6_computed_data.xlsx"),
    ("uvis", "UVIS_P6_computed_data.xlsx"),
    ("nir", "NIR_P6_computed_data.xlsx"),
)
SPECTRAL = tuple(
    f"{instrument}_P6_{kind}.xlsx"
    for instrument in ("UVIS", "NIR")
    for kind in ("absorbance", "transmitance", "scope")
)
EXCEL_ERRORS = {
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#N/A",
    "#NUM!",
    "#NAME?",
    "#NULL!",
    "#SPILL!",
    "#CALC!",
}
SESSION_FIELDS = (
    "session",
    *KEY_FIELDS,
    "id_status",
    "rgb_files",
    "noir_files",
    "rgb_container_ok",
    "noir_container_ok",
    "paired_view_indices",
    "exif_dates",
    "exposure_values",
    "iso_values",
)
IMAGE_FIELDS = (
    "path",
    "session",
    "camera",
    "view",
    "bytes",
    "width",
    "height",
    "datetime_original",
    "exposure_s",
    "iso",
    "container_check_error",
    "full_decode_ok",
)


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _csv(path: Path, rows: list[dict], fields: tuple | list) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8"
    )


def _integer(value: Any, context: str) -> int:
    try:
        result = int(value)
        if isinstance(value, bool) or float(value) != result or result < 1:
            raise ValueError
    except (TypeError, ValueError, OverflowError):
        raise ValueError(f"Invalid integer identity in {context}") from None
    return result


def _genotypes(rows: list[tuple]) -> tuple[dict[str, str], dict[str, int]]:
    lookup: dict[str, str] = {}
    ranks: dict[str, int] = {}
    for row in rows:
        if not any(value is not None for value in row):
            continue
        if len(row) != 2 or any(value is None or not str(value).strip() for value in row):
            raise ValueError("Incomplete genotype lookup row")
        name, code = (str(value).strip() for value in row)
        if code in lookup or name in ranks:
            raise ValueError("Duplicate or conflicting genotype lookup")
        lookup[code] = name
        ranks[name] = len(ranks)
    if not lookup:
        raise ValueError("Empty genotype lookup")
    return lookup, ranks


def _key(values: tuple | list, lookup: dict[str, str], context: str) -> tuple[str, int, str, int]:
    if len(values) < 4 or any(value is None or not str(value).strip() for value in values[:4]):
        raise ValueError(f"Incomplete identity key in {context}")
    treatment, measurement, genotype, replicate = values[:4]
    genotype = str(genotype).strip()
    genotype = lookup.get(genotype, genotype)
    if genotype not in lookup.values():
        raise ValueError(f"Unknown genotype identity in {context}")
    return (
        str(treatment).strip(),
        _integer(measurement, context),
        genotype,
        _integer(replicate, context),
    )


def _primary(rows: list[tuple], lookup: dict[str, str], name: str) -> tuple[list, dict, dict]:
    if not rows or [str(value).strip().lower() for value in rows[0][:4]] != [
        "treatment",
        "measurement",
        "gen",
        "rep",
    ]:
        raise ValueError(f"Unexpected four-key header in {name}")
    headers = list(rows[0])
    if len(set(str(header) for header in headers)) != len(headers) or any(
        header is None for header in headers
    ):
        raise ValueError(f"Missing or duplicate column header in {name}")
    indexed, coordinates = {}, {}
    for row_number, row in enumerate(rows[1:], 2):
        if all(value is None for value in row):
            continue
        key = _key(row, lookup, f"{name} row {row_number}")
        if key in indexed:
            raise ValueError(f"Duplicate identity key in {name} row {row_number}")
        for column, value in enumerate(row[4:], 5):
            cell = f"{get_column_letter(column)}{row_number}"
            if isinstance(value, str) and value in EXCEL_ERRORS:
                raise ValueError(f"Excel error in {name} {cell}")
            if (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(value)
            ):
                raise ValueError(f"Missing or nonnumeric measurement in {name} {cell}")
        indexed[key] = row
        coordinates[key] = row_number
    if not indexed:
        raise ValueError(f"No keyed records in {name}")
    return headers, indexed, coordinates


def _spectral_checks(tables: dict, keys: list[tuple]) -> tuple[list[dict], list[dict]]:
    issues, summaries = [], []
    for instrument in ("UVIS", "NIR"):
        absorbance, transmitted = (
            tables[f"{instrument}_P6_{kind}.xlsx"] for kind in ("absorbance", "transmitance")
        )
        headers, by_key, _ = absorbance
        max_error, invalid = 0.0, 0
        for key in keys:
            arow, trow = by_key[key], transmitted[1][key]
            for column, (a, t) in enumerate(zip(arow[4:], trow[4:]), 5):
                error = abs(a + math.log10(t)) if t > 0 else None
                if error is None:
                    invalid += 1
                else:
                    max_error = max(max_error, error)
                if instrument == "UVIS" and (error is None or error > 1e-5):
                    issues.append(
                        {
                            "key": list(key),
                            "wavelength": headers[column - 1],
                            "absorbance": a,
                            "transmittance": t,
                            "error": error,
                            "cell": f"{get_column_letter(column)}{transmitted[2][key]}",
                            "reason": "nonpositive_transmittance"
                            if error is None
                            else "absorbance_transmittance_disagreement",
                        }
                    )
        summaries.append(
            {
                "instrument_label": instrument,
                "wavelength_min": headers[4],
                "wavelength_max": headers[-1],
                "channels": len(headers) - 4,
                "A_vs_minus_log10_T_max_abs_error": max_error,
                "invalid_pairs": invalid,
            }
        )
    return issues, summaries


def _input_path(workspace: Path, relative: str) -> Path:
    logical = Path(str(relative).replace("\\", "/"))
    if logical.is_absolute() or PureWindowsPath(str(relative)).drive or ".." in logical.parts:
        raise ValueError("Source folder escapes outside the workspace")
    lexical = workspace / logical
    resolved = lexical.resolve()
    if resolved.is_relative_to(workspace.resolve()):
        return lexical
    registry = workspace / "input_bindings.json"
    bindings = json.loads(registry.read_text(encoding="utf-8-sig")) if registry.is_file() else {}
    for binding in bindings.get("bindings", []):
        prefix = Path(str(binding["logical_path"]).replace("\\", "/"))
        if prefix.is_absolute() or PureWindowsPath(str(prefix)).drive or ".." in prefix.parts:
            raise ValueError("Invalid registered input binding")
        if logical.is_relative_to(prefix):
            source = Path(binding["source_path"]).resolve()
            if resolved == (
                source / logical.relative_to(prefix)
            ).resolve() and resolved.is_relative_to(source):
                return lexical
    raise ValueError("Source path escapes registered input bindings")


def _session_inventory(
    workspace: Path, keys: set[tuple], lookup: dict[str, str]
) -> tuple[list[dict], list[dict], dict, dict]:
    config_path = workspace / "prototype/config.json"
    config = json.loads(config_path.read_text(encoding="utf-8")) if config_path.exists() else {}
    paths = set()
    for folder in config.get("source_folders", []):
        # Configured folders may be explicit registered bindings to source data.
        parent = _input_path(workspace, folder)
        if parent.is_dir():
            paths.update(path for path in parent.rglob("capture_session*") if path.is_dir())
    sessions, images, by_key = [], [], {}
    for session in sorted(paths):
        relative = session.relative_to(workspace).as_posix()
        _input_path(workspace, relative)
        match = re.fullmatch(r"capture_session_([CSP])(\d+)_(\d+)-(\d+)", session.name)
        key = None
        if match:
            treatment, measurement, genotype, replicate = match.groups()
            code = "L" + genotype
            if code not in lookup:
                raise ValueError(f"Unknown genotype code in session {relative}")
            key = _key((treatment, measurement, lookup[code], replicate), lookup, relative)
            if key in by_key:
                raise ValueError(f"Duplicate session identity: {relative}")
        record = {
            "session": relative,
            **dict(zip(KEY_FIELDS, key or ("", "", "", ""))),
            "id_status": "filename_key_matches_reference" if key in keys else "unresolved_or_pilot",
        }
        indices, dates, exposures, isos = {}, set(), set(), set()
        for camera in ("rgb", "noir"):
            files = sorted((session / camera).glob("*.jpg"))
            record[camera + "_files"] = len(files)
            record[camera + "_container_ok"] = 0
            indices[camera] = set()
            for path in files:
                _input_path(workspace, path.relative_to(workspace).as_posix())
                try:
                    view = int(path.stem.rsplit("_", 1)[1])
                except ValueError:
                    raise ValueError(
                        f"Unrecognized camera view index: {path.relative_to(workspace).as_posix()}"
                    ) from None
                if view in indices[camera]:
                    raise ValueError(f"Duplicate camera view index in {relative}/{camera}")
                indices[camera].add(view)
                width = height = date = exposure = iso = None
                error = ""
                try:
                    with Image.open(path) as image:
                        width, height = image.size
                        exif = image.getexif()
                        details = exif.get_ifd(34665) if 34665 in exif else {}
                        date = details.get(36867, exif.get(36867))
                        raw_exposure = details.get(33434, exif.get(33434))
                        exposure = float(raw_exposure) if raw_exposure is not None else None
                        iso = details.get(34855, exif.get(34855))
                        # JPEG verify() checks its header without decoding the
                        # entropy stream. Full load catches truncated pixels.
                        image.load()
                    record[camera + "_container_ok"] += 1
                except Exception as exc:
                    error = type(exc).__name__
                if date:
                    dates.add(str(date)[:10])
                if exposure is not None:
                    exposures.add(exposure)
                if iso is not None:
                    isos.add(str(iso))
                images.append(
                    {
                        "path": path.relative_to(workspace).as_posix(),
                        "session": relative,
                        "camera": camera,
                        "view": view,
                        "bytes": path.stat().st_size,
                        "width": width,
                        "height": height,
                        "datetime_original": date,
                        "exposure_s": exposure,
                        "iso": iso,
                        "container_check_error": error,
                        "full_decode_ok": not bool(error),
                    }
                )
        record.update(
            paired_view_indices=len(indices["rgb"] & indices["noir"]),
            exif_dates=";".join(sorted(dates)),
            exposure_values=";".join(map(str, sorted(exposures))),
            iso_values=";".join(sorted(isos)),
        )
        sessions.append(record)
        if key is not None:
            by_key[key] = record
    return sessions, images, by_key, config


def build_reference_tables(workspace: Path) -> dict[str, int]:
    """Export original sheets and validate a four-key OJIP/UVIS/NIR join.

    LAB and biomass exports are retained separately; their sampling linkage is
    unverified. No source workbook is modified and no key is joined by row order.
    """
    workspace = Path(workspace).absolute()
    paths = {path.name: path for path in sorted(workspace.glob("*.xlsx"))}
    required = {"Genotype_list.xlsx", *(name for _, name in COMPUTED), *SPECTRAL}
    missing = required - paths.keys()
    if missing:
        raise ValueError("Missing source workbooks: " + ", ".join(sorted(missing)))
    books = {}
    hashes = {}
    for name, path in paths.items():
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            books[name] = [(sheet.title, list(sheet.values)) for sheet in book]
        finally:
            book.close()
        hashes[name] = {"bytes": path.stat().st_size, "sha256": _sha(path)}
    lookup, genotype_rank = _genotypes(books["Genotype_list.xlsx"][0][1])
    tables = {}
    for name, sheets in books.items():
        if name in required - {"Genotype_list.xlsx"} or name == "OJIP_P6_source_data.xlsx":
            tables[name] = _primary(sheets[0][1], lookup, name)
    anchor = tables[COMPUTED[0][1]][1]
    keys = sorted(anchor, key=lambda key: (key[0], key[1], genotype_rank[key[2]], key[3]))
    keyset = set(keys)
    alignments = []
    for name, (headers, rows, _) in tables.items():
        if set(rows) != keyset:
            raise ValueError(
                f"Key sets differ from OJIP in {name}: missing={len(keyset - set(rows))}, extra={len(set(rows) - keyset)}"
            )
        alignments.append(
            {
                "file": name,
                "matched_ojip_keys": len(keyset),
                "missing_vs_ojip": 0,
                "extra_vs_ojip": 0,
            }
        )
        if name in SPECTRAL:
            grid = headers[4:]
            if len(grid) != 256 or any(
                isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v)
                for v in grid
            ):
                raise ValueError(
                    f"Spectral wavelength grid must contain 256 finite native channels in {name}"
                )
            if any(a >= b for a, b in zip(grid, grid[1:])):
                raise ValueError(f"Wavelength grid must be strictly increasing in {name}")
    for instrument in ("UVIS", "NIR"):
        grids = [
            tables[f"{instrument}_P6_{kind}.xlsx"][0][4:]
            for kind in ("absorbance", "transmitance", "scope")
        ]
        if any(grid != grids[0] for grid in grids[1:]):
            raise ValueError(f"Wavelength grids differ within {instrument}")
    issues, spectral = _spectral_checks(tables, keys)
    sessions, images, session_by_key, config = _session_inventory(workspace, keyset, lookup)
    master, dictionary = [], []
    fields = list(KEY_FIELDS)
    for prefix, name in COMPUTED:
        headers = tables[name][0]
        fields.extend(f"{prefix}__{header}" for header in headers[4:])
        dictionary.extend(
            {
                "file": name,
                "sheet": books[name][0][0],
                "column": get_column_letter(column),
                "header_verbatim": header,
                "meaning_status": "identifier"
                if column <= 4
                else "definition_and_units_require_protocol_verification",
            }
            for column, header in enumerate(headers, 1)
        )
    fields.extend(("image_session", "image_status"))
    allowed = config.get("allowed_views_per_camera", [30, 40])
    for key in keys:
        row = dict(zip(KEY_FIELDS, key))
        for prefix, name in COMPUTED:
            headers, values, _ = tables[name]
            row.update(
                (f"{prefix}__{header}", value)
                for header, value in zip(headers[4:], values[key][4:])
            )
        session = session_by_key.get(key)
        row.update(
            image_session=session["session"] if session else "",
            image_status="no_identified_extracted_session",
        )
        if session:
            count = session["rgb_files"]
            complete = count in allowed and all(
                session[field] == count
                for field in (
                    "noir_files",
                    "rgb_container_ok",
                    "noir_container_ok",
                    "paired_view_indices",
                )
            )
            row["image_status"] = (
                f"complete_{count}_plus_{count}_container_checks"
                if complete
                else "incomplete_extracted_session"
            )
        master.append(row)
    # Validate every input before publishing derived tables.
    audit = workspace / "research_audit"
    audit.mkdir(parents=True, exist_ok=True)
    for name, sheets in books.items():
        for title, rows in sheets:
            target = audit / "tables" / f"{Path(name).stem}__{title}.csv"
            if Path(title).name != title or title in (".", ".."):
                raise ValueError(f"Unsafe sheet title in {name}")
            target.parent.mkdir(parents=True, exist_ok=True)
            with target.open("w", encoding="utf-8-sig", newline="") as stream:
                csv.writer(stream).writerows(rows)
    _csv(audit / "reference_master.csv", master, fields)
    _csv(
        audit / "computed_variable_dictionary.csv",
        dictionary,
        ("file", "sheet", "column", "header_verbatim", "meaning_status"),
    )
    _csv(
        audit / "reference_key_alignment.csv",
        alignments,
        ("file", "matched_ojip_keys", "missing_vs_ojip", "extra_vs_ojip"),
    )
    _csv(
        audit / "reference_workbook_sha256.csv",
        [{"path": name, **entry} for name, entry in hashes.items()],
        ("path", "bytes", "sha256"),
    )
    _csv(audit / "session_manifest.csv", sessions, SESSION_FIELDS)
    _csv(audit / "image_manifest.csv", images, IMAGE_FIELDS)
    _csv(
        audit / "spectral_consistency.csv",
        spectral,
        (
            "instrument_label",
            "wavelength_min",
            "wavelength_max",
            "channels",
            "A_vs_minus_log10_T_max_abs_error",
            "invalid_pairs",
        ),
    )
    _json(audit / "uvis_consistency_issues.json", issues)
    counts = {
        "reference_records": len(master),
        "computed_variables": len(dictionary),
        "source_workbooks": len(books),
        "exported_sheets": sum(len(sheets) for sheets in books.values()),
        "uvis_consistency_issues": len(issues),
        "sessions": len(sessions),
        "images": len(images),
    }
    sources = {
        "schema_version": 1,
        "join_keys": list(KEY_FIELDS),
        "workbooks": hashes,
        "builder_sha256": _sha(Path(__file__)),
        "openpyxl_version": openpyxl.__version__,
        "computed_sources": dict(COMPUTED),
        "spectral_channels_per_instrument": 256,
        "assay_join_status": "not_joined_sampling_relationship_unverified",
        "spectral_consistency_tolerance": 1e-5,
        "source_folder_policy": "Only configured source_folders; timestamp identities remain unresolved.",
        "counts": counts,
    }
    if (workspace / "prototype/config.json").exists():
        sources["config_sha256"] = _sha(workspace / "prototype/config.json")
    _json(workspace / "metadata/reference_sources.json", sources)
    return counts
