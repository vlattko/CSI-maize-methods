"""Workbook ingestion contracts using synthetic Excel acquisitions."""

import csv
import hashlib
import importlib
import json
from pathlib import Path
import tempfile
import unittest

import openpyxl
from PIL import Image

KEY_HEADERS = ["Treatment", "Measurement", "GEN", "REP"]
KEY_A = ["C", 1, "Alpha", 1]
KEY_B = ["C", 1, "Beta", 1]


def workbook(path, rows, title="List1"):
    book = openpyxl.Workbook()
    book.active.title = title
    for row in rows:
        book.active.append(row)
    book.save(path)
    book.close()


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


class MetadataTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.audit = self.root / "research_audit"
        workbook(self.root / "Genotype_list.xlsx", [["Alpha", "L1"], ["Beta", "L2"]])
        workbook(
            self.root / "OJIP_P6_computed_data.xlsx",
            [KEY_HEADERS + ["Fv/Fm", "Pi_Abs"], KEY_A + [0.8, 2], KEY_B + [0.7, 3]],
            "SVE",
        )
        workbook(
            self.root / "UVIS_P6_computed_data.xlsx",
            [KEY_HEADERS + ["NDVI"], KEY_B + [0.9], KEY_A + [0.3]],
        )
        workbook(
            self.root / "NIR_P6_computed_data.xlsx",
            [KEY_HEADERS + ["NDVI"], KEY_B + [0.2], KEY_A + [0.1]],
        )
        for instrument, start in [("UVIS", 325), ("NIR", 628)]:
            for kind, value in [("absorbance", 1), ("transmitance", 0.1), ("scope", 10)]:
                workbook(
                    self.root / f"{instrument}_P6_{kind}.xlsx",
                    [
                        KEY_HEADERS + list(range(start, start + 256)),
                        KEY_B + [value] * 256,
                        KEY_A + [value] * 256,
                    ],
                )
        workbook(
            self.root / "LAB_1.xlsx",
            [["TREATMENT", "GEN", None, "chlorophyll"], ["C", "L1", 1, 99]],
        )
        workbook(self.root / "MASA_P_6.xlsx", [[None, None, "M_C_fresh"], ["L1", "Alpha", 5]])

    def build(self):
        try:
            module = importlib.import_module("csi_maize.metadata")
        except ModuleNotFoundError as error:
            self.fail(f"Workbook ingestion is not implemented: {error.name}")
        return module.build_reference_tables(self.root)

    def change(self, name, edit):
        path = self.root / name
        book = openpyxl.load_workbook(path)
        edit(book.active)
        book.save(path)
        book.close()

    def test_joins_four_keys_instead_of_row_positions_without_inventing_assay_links(self):
        result = self.build()
        rows = read_csv(self.audit / "reference_master.csv")
        self.assertEqual(result["reference_records"], 2)
        self.assertEqual(
            [(r["genotype"], r["ojip__Fv/Fm"], r["uvis__NDVI"], r["nir__NDVI"]) for r in rows],
            [("Alpha", "0.8", "0.3", "0.1"), ("Beta", "0.7", "0.9", "0.2")],
        )
        self.assertFalse(
            any(
                "lab" in key.lower() or "chlorophyll" in key.lower() or "fresh" in key
                for key in rows[0]
            )
        )
        self.assertEqual(rows[0]["image_status"], "no_identified_extracted_session")
        self.assertTrue((self.audit / "tables/LAB_1__List1.csv").is_file())
        sources = json.loads((self.root / "metadata/reference_sources.json").read_text())
        self.assertEqual(
            sources["assay_join_status"], "not_joined_sampling_relationship_unverified"
        )
        original = self.root / "OJIP_P6_computed_data.xlsx"
        self.assertEqual(
            sources["workbooks"][original.name]["sha256"],
            hashlib.sha256(original.read_bytes()).hexdigest(),
        )
        self.assertNotIn(str(self.root), json.dumps(sources))

    def test_row_permutation_does_not_change_reference_values_or_order(self):
        self.build()
        expected = (self.audit / "reference_master.csv").read_bytes()
        for path in self.root.glob("*_P6_*.xlsx"):
            if path.name.startswith("MASA"):
                continue

            def reverse(sheet):
                rows = list(sheet.values)
                sheet.delete_rows(2, sheet.max_row)
                for row in reversed(rows[1:]):
                    sheet.append(row)

            self.change(path.name, reverse)
        self.build()
        self.assertEqual((self.audit / "reference_master.csv").read_bytes(), expected)

    def test_duplicate_key_is_rejected_even_when_values_agree(self):
        self.change("UVIS_P6_computed_data.xlsx", lambda s: s.append(KEY_A + [0.3]))
        with self.assertRaisesRegex(ValueError, "[Dd]uplicate.*key"):
            self.build()
        self.assertFalse((self.audit / "reference_master.csv").exists())

    def test_missing_key_is_not_silently_inner_joined(self):
        self.change("NIR_P6_computed_data.xlsx", lambda s: s.delete_rows(2))
        with self.assertRaisesRegex(ValueError, "[Kk]ey.*(?:differ|mismatch)|[Mm]issing.*key"):
            self.build()

    def test_incomplete_key_is_rejected(self):
        self.change("UVIS_P6_computed_data.xlsx", lambda s: setattr(s["C2"], "value", None))
        with self.assertRaisesRegex(ValueError, "[Ii]ncomplete|[Mm]issing.*identity"):
            self.build()

    def test_excel_error_is_not_coerced_to_a_valid_numeric_measurement(self):
        self.change("OJIP_P6_computed_data.xlsx", lambda s: setattr(s["E2"], "value", "#DIV/0!"))
        with self.assertRaisesRegex(ValueError, "Excel error.*E2"):
            self.build()

    def test_genotype_lookup_cannot_assign_one_code_to_two_names(self):
        self.change("Genotype_list.xlsx", lambda s: s.append(["Gamma", "L1"]))
        with self.assertRaisesRegex(ValueError, "[Cc]onflicting.*genotype|[Dd]uplicate.*genotype"):
            self.build()

    def test_native_spectral_grid_requires_256_channels_and_matching_wavelengths(self):
        for break_grid in [
            lambda s: s.delete_cols(260),
            lambda s: setattr(s["E1"], "value", 324.5),
        ]:
            with self.subTest(break_grid=break_grid):
                self.change("UVIS_P6_transmitance.xlsx", break_grid)
                with self.assertRaisesRegex(ValueError, "256|[Ww]avelength"):
                    self.build()
                workbook(
                    self.root / "UVIS_P6_transmitance.xlsx",
                    [KEY_HEADERS + list(range(325, 581)), KEY_B + [0.1] * 256, KEY_A + [0.1] * 256],
                )

    def test_spectral_flags_preserve_source_coordinates_and_use_strict_json(self):
        self.change("UVIS_P6_transmitance.xlsx", lambda s: setattr(s["E2"], "value", -0.1))
        self.change("UVIS_P6_transmitance.xlsx", lambda s: setattr(s["F3"], "value", 1))
        self.build()
        raw = (self.audit / "uvis_consistency_issues.json").read_text()
        issues = json.loads(
            raw, parse_constant=lambda value: self.fail(f"Nonstandard JSON: {value}")
        )
        negative = next(x for x in issues if x["cell"] == "E2")
        disagreement = next(x for x in issues if x["cell"] == "F3")
        self.assertEqual(negative["key"], KEY_B)
        self.assertEqual(negative["wavelength"], 325)
        self.assertIsNone(negative["error"])
        self.assertEqual(negative["reason"], "nonpositive_transmittance")
        self.assertEqual(disagreement["key"], KEY_A)
        self.assertEqual(disagreement["error"], 1)

    def test_session_manifest_comes_from_images_and_timestamp_names_remain_unresolved(self):
        folder = self.root / "images"
        named = folder / "capture_session_C1_1-1"
        timestamp = folder / "capture_session_2025-03-26_15-29-25"
        for session in [named, timestamp]:
            for camera in ["rgb", "noir"]:
                (session / camera).mkdir(parents=True)
                Image.new("RGB", (8, 6), "green").save(session / camera / "view_0.jpg")
        (self.root / "prototype").mkdir()
        (self.root / "prototype/config.json").write_text(
            json.dumps({"source_folders": ["images"], "allowed_views_per_camera": [30, 40]})
        )
        self.build()
        sessions = read_csv(self.audit / "session_manifest.csv")
        named_row = next(r for r in sessions if r["genotype"])
        unknown = next(r for r in sessions if not r["genotype"])
        self.assertEqual(named_row["genotype"], "Alpha")
        self.assertEqual(named_row["paired_view_indices"], "1")
        self.assertEqual(named_row["rgb_container_ok"], "1")
        self.assertEqual(unknown["id_status"], "unresolved_or_pilot")
        master = read_csv(self.audit / "reference_master.csv")
        self.assertEqual(master[0]["image_session"], "images/capture_session_C1_1-1")
        self.assertEqual(master[0]["image_status"], "incomplete_extracted_session")
        self.assertEqual(master[1]["image_session"], "")

    def test_only_registered_data_links_may_escape_the_workspace(self):
        with tempfile.TemporaryDirectory() as raw_directory:
            source = Path(raw_directory)
            (source / "capture_session_C1_1-1").mkdir()
            try:
                (self.root / "images").symlink_to(source, target_is_directory=True)
            except OSError:
                self.skipTest("Directory symlinks are unavailable")
            (self.root / "prototype").mkdir()
            (self.root / "prototype/config.json").write_text(
                json.dumps({"source_folders": ["images"]})
            )
            with self.assertRaisesRegex(ValueError, "registered|binding"):
                self.build()
            (self.root / "input_bindings.json").write_text(
                json.dumps({"bindings": [{"logical_path": "images", "source_path": str(source)}]})
            )
            self.build()
            self.assertEqual(read_csv(self.audit / "session_manifest.csv")[0]["genotype"], "Alpha")

    def test_complete_container_counts_do_not_hide_a_damaged_camera_view(self):
        session = self.root / "images/capture_session_C1_1-1"
        for camera in ["rgb", "noir"]:
            (session / camera).mkdir(parents=True)
            for view in range(40):
                Image.new("RGB", (8, 6), "green").save(session / camera / f"view_{view}.jpg")
        (self.root / "prototype").mkdir()
        (self.root / "prototype/config.json").write_text(json.dumps({"source_folders": ["images"]}))
        self.build()
        self.assertEqual(
            read_csv(self.audit / "reference_master.csv")[0]["image_status"],
            "complete_40_plus_40_container_checks",
        )
        (session / "noir/view_0.jpg").write_bytes(b"damaged")
        self.build()
        self.assertEqual(
            read_csv(self.audit / "reference_master.csv")[0]["image_status"],
            "incomplete_extracted_session",
        )
        self.assertEqual(
            read_csv(self.audit / "session_manifest.csv")[0]["noir_container_ok"], "39"
        )

    def test_duplicate_session_identity_and_source_path_escape_are_rejected(self):
        for directory in ["a", "b"]:
            (self.root / "images" / directory / "capture_session_C1_1-1").mkdir(parents=True)
        (self.root / "prototype").mkdir()
        config = self.root / "prototype/config.json"
        config.write_text(json.dumps({"source_folders": ["images"]}))
        with self.assertRaisesRegex(ValueError, "[Dd]uplicate.*session|[Mm]ultiple.*session"):
            self.build()
        config.write_text(json.dumps({"source_folders": ["../outside"]}))
        with self.assertRaisesRegex(ValueError, "outside|escape"):
            self.build()


if __name__ == "__main__":
    unittest.main()
