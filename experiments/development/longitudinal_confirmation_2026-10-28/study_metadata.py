"""Read-only timetable and experimenter-confirmed destructive-mass crosswalk."""
import argparse
import csv
import datetime as dt
import html
import json
import math
from pathlib import Path

from registry import sha


AGES = {13: '1', 19: '2', 20: '3', 22: '4', 25: '5'}


def mass_records(rows, genotype_codes, confirmation):
    if (confirmation.get('replicate_blocks_confirmed') is not True
            or confirmation.get('mass_units') != 'g'
            or confirmation.get('tissue') != 'whole aboveground plant'):
        raise ValueError('Explicit block, unit and tissue confirmation required')
    if len(genotype_codes) != 6 or len(rows) != 24:
        raise ValueError('Expected four six-genotype blocks')
    result = []
    for block in range(4):
        group = rows[6*block:6*(block+1)]
        if {r[1] for r in group} != set(genotype_codes):
            raise ValueError('Each replicate block must contain six distinct named genotypes')
        for offset, row in enumerate(group):
            source_row = 2 + 6*block + offset
            if len(row) != 8 or genotype_codes[row[1]] != row[0]:
                raise ValueError('Genotype name/code disagreement or unexpected columns')
            for treatment, column, letters in [('S', 2, 'CD'), ('C', 4, 'EF'), ('P', 6, 'GH')]:
                fresh, dry = row[column:column+2]
                if (not all(isinstance(v, (int, float)) and not isinstance(v, bool)
                            and math.isfinite(v) for v in (fresh, dry))
                        or not 0 < dry <= fresh):
                    raise ValueError('Fresh/dry mass must be finite and 0 < dry <= fresh')
                result.append(dict(treatment=treatment, genotype=row[1], replicate=str(block+1),
                                   fresh_mass_g=float(fresh), dry_mass_g=float(dry),
                                   dry_matter_fraction=dry/fresh, water_mass_g=fresh-dry,
                                   source_sheet='List1', source_fresh_cell=f'{letters[0]}{source_row}',
                                   source_dry_cell=f'{letters[1]}{source_row}',
                                   tissue=confirmation['tissue'], mass_units='g',
                                   mapping_authority='direct experimenter confirmation'))
    return result

def timetable_records(rows):
    blocks = []
    for row in rows:
        if not isinstance(row.get('das'), (int, float)):
            continue
        if row['das'] != int(row['das']) or row['das'] < 0:
            raise ValueError('Invalid scheduled DAS')
        if row['das'] == 0:
            blocks.append([])
        if not blocks:
            raise ValueError('Scheduled rows must start with sowing')
        blocks[-1].append(row)
    result = []
    seen = set()
    for block in blocks:
        text = ' '.join(str(r['action']) for r in block).lower()
        treatment = 'P' if 'priming' in text else ('S' if 'cold' in text else 'C')
        if treatment in seen:
            raise ValueError('Ambiguous repeated treatment block')
        seen.add(treatment)
        start = block[0]['date']
        ages = set()
        for row in block:
            age = int(row['das'])
            if not isinstance(row['date'], (dt.datetime, dt.date)) or (row['date']-start).days != age or age in ages:
                raise ValueError('Scheduled date/DAS mismatch or duplicate age')
            ages.add(age)
            result.append(dict(treatment=treatment, scheduled_das=age,
                               scheduled_date=row['date'].date().isoformat() if isinstance(row['date'], dt.datetime) else row['date'].isoformat(),
                               measurement=AGES.get(age, '') if 'snimanje' in str(row['detail']).lower() else '',
                               action=row['action'], detail=row['detail'],
                               day_temperature_raw=row['day'], night_temperature_raw=row['night'],
                               source_sheet='List1', source_row=row['row'],
                               treatment_mapping='inferred from priming/cold actions in sowing block',
                               date_status='scheduled date; not an image timestamp'))
    return result


def write_csv(path, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ['timetable', 'mass', 'genotypes', 'mass-confirmation', 'output']:
        parser.add_argument('--'+name, required=True, type=Path)
    args = parser.parse_args()
    import openpyxl
    args.output.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(args.timetable, data_only=True)
    sheet = workbook['List1']
    raw = []
    cells = []
    for row in sheet:
        for cell in row:
            if cell.value is not None or cell.comment is not None:
                cells.append(dict(cell=cell.coordinate, value=cell.value,
                                  comment=cell.comment.text if cell.comment else None))
        values = [sheet.cell(row[0].row, col).value for col in range(1, 8)]
        raw.append(dict(row=row[0].row, date=values[1], das=values[4], action=values[2],
                        detail=values[3], day=values[5], night=values[6]))
    schedule = timetable_records(raw)
    if len(schedule) != 84 or {r['treatment'] for r in schedule} != {'C', 'P', 'S'}:
        raise ValueError('Expected three 28-day schedule blocks')
    visits = [r for r in schedule if r['measurement']]
    if len(visits) != 15 or any(sorted(int(r['scheduled_das']) for r in visits if r['treatment']==t) != sorted(AGES) for t in ['C','P','S']):
        raise ValueError('Incomplete scheduled acquisition grid')
    genotype_rows = list(openpyxl.load_workbook(args.genotypes, data_only=True).active.values)
    codes = dict(genotype_rows)
    if len(codes) != len(genotype_rows):
        raise ValueError('Duplicate genotype in crosswalk')
    mass_sheet = openpyxl.load_workbook(args.mass, data_only=True)['List1']
    mass_rows = list(mass_sheet.values)
    if tuple(mass_rows[0][2:]) != ('M_S_fresh','M_S_dry','M_C_fresh','M_C_dry','M_P_fresh','M_P_dry'):
        raise ValueError('Unexpected mass headers')
    confirmation = json.loads(args.mass_confirmation.read_text())
    masses = mass_records(mass_rows[1:], codes, confirmation)
    write_csv(args.output/'scheduled_days.csv', schedule)
    write_csv(args.output/'scheduled_visits.csv', visits)
    write_csv(args.output/'confirmed_mass.csv', masses)
    reconciliation = dict(
        cold_challenge=dict(treatments=['S','P'], start_DAS=14, end_DAS=19,
                            day_C=10, night_C=5, last_day='25 C restored after measurements',
                            exact_hours=None),
        priming=dict(start_DAS=6, end_DAS=11, day_C=12, working_night_C=8,
                     duration_pairs_raw=['4+2','6+3','8+4','10+5','10+5','10+5'],
                     interpretation='Increasing exposure duration; day/night hours inferred, awaiting complete method',
                     night_conflict=dict(cells='G80:G85: 8/18**', footnote='B105: 5 C',
                                         experimenter='8 C', status='unresolved; direct experimenter value used provisionally')),
        first_imaging='13 DAS: before severe challenge, but after priming for P',
        recovery_imaging_DAS=[20,22,25], LAB_DAS=[19,25], fresh_mass_DAS=26, dry_matter_schedule_DAS=27,
        constraints=['Scheduled dates differ between treatments; treatment and calendar run may be confounded.',
                     'Scheduled dates do not replace image timestamps; actual execution dates require reconciliation.',
                     'LAB_1/LAB_2 labels are not automatically assigned to these scheduled LAB events.',
                     'LAB tissue removal is unknown; missing later leaf regions do not establish senescence.',
                     'DAS 19 cold-to-control transition follows measurement; six labelled days do not imply 144 cold hours.'])
    document = dict(source_sha256=sha(args.timetable), sheet='List1', cells=cells,
                    reconciliation=reconciliation, mass_confirmation=confirmation)
    (args.output/'metadata_reconciliation.json').write_text(json.dumps(document, indent=2, default=str), encoding='utf-8')
    head = '<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Study schedule and mass mapping</title><style>body{font:19px/1.5 "Times New Roman",serif;max-width:1200px;margin:35px auto;padding:24px;background:white;color:#171717}table{border-collapse:collapse;width:100%;font-size:17px}td,th{border-bottom:1px solid #ddd;text-align:left;padding:7px}pre{white-space:pre-wrap}svg{width:100%;height:auto}</style><h1>Study schedule and destructive mass measurements</h1><p>The timetable supplies 84 scheduled treatment-days and 15 acquisition visits. The experimenter confirms the mass units as grams, the tissue as the whole aboveground plant and the four six-genotype blocks as successive replicates. We retain the written genotype within every block, whose internal order varies.</p>'
    svg = '<svg viewBox="0 0 1000 210" role="img" aria-label="Treatment schedule on DAS axis">'
    colors = {'control':'#deded8','cold':'#746395','priming':'#b28a47'}
    for i,t in enumerate(['C','S','P']):
        y=20+55*i
        svg += f'<text x="0" y="{y+23}">{t}</text>'
        for d in range(28):
            stage='cold' if t in ['S','P'] and 14<=d<=19 else ('priming' if t=='P' and 6<=d<=11 else 'control')
            svg += f'<rect x="{40+d*33}" y="{y}" width="32" height="28" fill="{colors[stage]}"><title>{t}, {d} DAS: {stage}; day 19 transitions after measurement</title></rect>'
            if d in AGES:svg += f'<circle cx="{56+d*33}" cy="{y+14}" r="4" fill="black"/>'
    for d in [0,6,11,13,14,19,20,22,25,26,27]:svg += f'<text x="{56+d*33}" y="202" text-anchor="middle" font-size="15">{d}</text>'
    svg += '</svg><p>DAS axis. Grey: control conditions; ochre: intermittent priming; purple: severe cold, with warming after the DAS 19 measurements; dots: imaging. P has already been primed at the first image.</p>'
    rows_html = ''.join('<tr>'+''.join('<td>'+html.escape(str(r[k] or ''))+'</td>' for k in ['treatment','measurement','scheduled_das','scheduled_date','detail'])+'</tr>' for r in visits)
    page = head+svg+'<h2>Acquisition timetable</h2><table><tr><th>Treatment</th><th>Visit</th><th>DAS</th><th>Scheduled date</th><th>Actions</th></tr>'+rows_html+'</table><h2>Interpretation and unresolved details</h2><pre>'+html.escape(json.dumps(reconciliation,indent=2))+'</pre><p>These are schedule records, not a replacement for camera timestamps. The mass-to-pot join is confirmed independently of camera geometry. The last imaging visit precedes scheduled fresh harvest by one day. No raw files were modified.</p></html>'
    (args.output/'study_schedule.html').write_text(page, encoding='utf-8')
    inputs=[args.timetable,args.mass,args.genotypes,args.mass_confirmation]
    manifest=dict(status='completed', scheduled_days=len(schedule), scheduled_visits=len(visits), confirmed_mass_pots=len(masses),
                  inputs={str(p):sha(p) for p in inputs}, source_sha256=sha(Path(__file__)),
                  outputs={p.name:sha(p) for p in args.output.iterdir() if p.is_file() and p.name!='metadata_manifest.json'})
    (args.output/'metadata_manifest.json').write_text(json.dumps(manifest,indent=2), encoding='utf-8')
    print(json.dumps({k:manifest[k] for k in ['status','scheduled_days','scheduled_visits','confirmed_mass_pots']}))


if __name__ == '__main__':
    main()
